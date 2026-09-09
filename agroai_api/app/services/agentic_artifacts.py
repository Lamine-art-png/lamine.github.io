from __future__ import annotations

import base64
import hashlib
import io
import os
import re
import tempfile
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from app.api.v1.chat_artifacts import (
    BRAND_GREEN,
    BRAND_LIME,
    BRAND_LOGO_PNG_BASE64,
    BRAND_MUTED,
    ReportPdfRequest,
    build_report_pdf_bytes,
)
from app.models.operational_records import GeneratedArtifact
from app.services.object_storage import get_object_store, object_storage_configured


_ARTIFACT_SCOPE = "agentic-artifacts"
_FORMATS = {
    "pdf": ("application/pdf", ".pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
    "pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", ".pptx"),
}


def _plain(value: Any, limit: int = 12000) -> str:
    return str(value or "").strip()[:limit]


def _filename(title: str, extension: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", title).strip("-").lower() or "agro-ai-artifact"
    return f"{slug[:80]}{extension}"


def _blocks(text: str, *, max_blocks: int = 10) -> list[str]:
    raw = _plain(text, 50000)
    chunks = [part.strip() for part in re.split(r"\n\s*\n|(?<=\.)\s+(?=[A-Z])", raw) if part.strip()]
    return chunks[:max_blocks] or ["No analysis content was supplied."]


def _evidence_lines(items: list[dict[str, Any]]) -> list[str]:
    lines: list[str] = []
    for item in (items or [])[:20]:
        name = str(item.get("filename") or item.get("name") or item.get("source_type") or "Workspace source")
        rows = item.get("rows_parsed") or item.get("rows")
        bits = [name]
        if rows is not None:
            bits.append(f"{rows} rows")
        warnings = item.get("warnings") or []
        if warnings:
            bits.append(f"{len(warnings)} warning(s)")
        lines.append(" · ".join(bits))
    return lines


def _docx_bytes(*, title: str, question: str, answer: str, evidence: list[dict[str, Any]]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    header = section.header.paragraphs[0]
    header.style = doc.styles["Normal"]
    try:
        logo_run = header.add_run()
        logo_run.add_picture(io.BytesIO(base64.b64decode(BRAND_LOGO_PNG_BASE64)), width=Inches(0.34))
        brand_run = header.add_run("  AGRO-AI")
    except Exception:
        brand_run = header.add_run("AGRO-AI")
    brand_run.bold = True
    brand_run.font.size = Pt(10)
    brand_run.font.color.rgb = RGBColor(13, 43, 30)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(13, 43, 30)

    sub = doc.add_paragraph("AGRO-AI workspace artifact")
    sub.runs[0].font.size = Pt(10)
    sub.runs[0].font.color.rgb = RGBColor(102, 116, 103)

    doc.add_heading("Request", level=1)
    doc.add_paragraph(question or "Workspace artifact requested through Ask AGRO-AI.")

    doc.add_heading("Executive summary", level=1)
    for block in _blocks(answer, max_blocks=7):
        doc.add_paragraph(block)

    doc.add_heading("Workspace evidence", level=1)
    evidence_lines = _evidence_lines(evidence)
    if evidence_lines:
        for line in evidence_lines:
            doc.add_paragraph(line, style="List Bullet")
    else:
        doc.add_paragraph("No uploaded evidence metadata was attached to this artifact.")

    doc.add_heading("Operating note", level=1)
    doc.add_paragraph(
        "This artifact is generated from the authenticated workspace context. "
        "Source provenance, missing data, and human approval boundaries remain authoritative."
    )

    styles = doc.styles
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.color.rgb = RGBColor(13, 43, 30)
    styles["Heading 1"].font.size = Pt(15)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _pptx_bytes(*, title: str, question: str, answer: str, evidence: list[dict[str, Any]]) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    green = RGBColor(13, 43, 30)
    lime = RGBColor(167, 224, 58)
    offwhite = RGBColor(248, 248, 244)
    muted = RGBColor(102, 116, 103)
    white = RGBColor(255, 255, 255)

    def add_text(slide, text, left, top, width, height, *, size=20, bold=False, color=green, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        p = frame.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return box

    logo_bytes = base64.b64decode(BRAND_LOGO_PNG_BASE64)

    def brand(slide, dark=False):
        try:
            slide.shapes.add_picture(io.BytesIO(logo_bytes), Inches(0.65), Inches(0.18), height=Inches(0.42))
            text_left = Inches(1.15)
        except Exception:
            text_left = Inches(0.65)
        add_text(
            slide,
            "AGRO-AI",
            text_left, Inches(0.28), Inches(2.0), Inches(0.35),
            size=10, bold=True, color=white if dark else green,
        )
        line = slide.shapes.add_shape(1, Inches(0.65), Inches(7.05), Inches(1.0), Inches(0.055))
        line.fill.solid()
        line.fill.fore_color.rgb = lime
        line.line.fill.background()

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = green
    brand(slide, dark=True)
    add_text(slide, title, Inches(0.9), Inches(1.65), Inches(11.5), Inches(2.0), size=34, bold=True, color=white)
    add_text(slide, question[:220], Inches(0.92), Inches(4.05), Inches(10.5), Inches(1.0), size=16, color=RGBColor(204, 218, 208))
    add_text(slide, "Workspace intelligence artifact", Inches(0.92), Inches(6.35), Inches(5.0), Inches(0.4), size=10, color=lime)

    blocks = _blocks(answer, max_blocks=8)

    # Summary
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = offwhite
    brand(slide)
    add_text(slide, "Executive summary", Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.6), size=26, bold=True)
    y = 1.75
    for block in blocks[:3]:
        add_text(slide, block[:520], Inches(0.9), Inches(y), Inches(11.3), Inches(1.15), size=15, color=RGBColor(38, 58, 48))
        y += 1.42

    # Findings
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = white
    brand(slide)
    add_text(slide, "What the operation shows", Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.6), size=26, bold=True)
    y = 1.75
    for index, block in enumerate(blocks[3:6] or blocks[:3], start=1):
        add_text(slide, f"{index:02d}", Inches(0.9), Inches(y), Inches(0.65), Inches(0.45), size=12, bold=True, color=muted)
        add_text(slide, block[:500], Inches(1.65), Inches(y - 0.02), Inches(10.5), Inches(1.0), size=15, color=RGBColor(38, 58, 48))
        y += 1.5

    # Evidence
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = offwhite
    brand(slide)
    add_text(slide, "Evidence & source coverage", Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.6), size=26, bold=True)
    evidence_lines = _evidence_lines(evidence)
    if not evidence_lines:
        evidence_lines = ["No uploaded evidence metadata attached to this request."]
    y = 1.75
    for line in evidence_lines[:6]:
        add_text(slide, "•", Inches(0.95), Inches(y), Inches(0.35), Inches(0.35), size=16, bold=True, color=lime)
        add_text(slide, line[:240], Inches(1.35), Inches(y - 0.02), Inches(10.7), Inches(0.55), size=14, color=RGBColor(38, 58, 48))
        y += 0.78

    # Action slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = green
    brand(slide, dark=True)
    add_text(slide, "Turn intelligence into action", Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.6), size=26, bold=True, color=white)
    action_blocks = blocks[6:] or [
        "Validate the highest-impact finding against its source evidence.",
        "Assign the next operational task to the responsible team member.",
        "Keep the resulting decision and evidence trail inside the workspace.",
    ]
    y = 1.9
    for index, block in enumerate(action_blocks[:4], start=1):
        add_text(slide, f"{index}.", Inches(0.95), Inches(y), Inches(0.5), Inches(0.45), size=16, bold=True, color=lime)
        add_text(slide, block[:460], Inches(1.55), Inches(y - 0.04), Inches(10.6), Inches(0.9), size=16, color=white)
        y += 1.15

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def build_artifact_bytes(
    *,
    format_name: str,
    title: str,
    question: str,
    answer: str,
    uploaded_evidence: list[dict[str, Any]],
    tenant_id: str,
) -> tuple[bytes, str, str]:
    fmt = str(format_name or "pdf").strip().lower()
    if fmt not in _FORMATS:
        raise ValueError("Unsupported artifact format")
    content_type, extension = _FORMATS[fmt]
    clean_title = _plain(title, 180) or "AGRO-AI Operating Artifact"
    if fmt == "pdf":
        data = build_report_pdf_bytes(
            ReportPdfRequest(
                title=clean_title,
                question=_plain(question, 12000),
                answer=_plain(answer, 50000),
                uploaded_evidence=uploaded_evidence or [],
            ),
            tenant_id,
        )
    elif fmt == "docx":
        data = _docx_bytes(
            title=clean_title,
            question=_plain(question, 12000),
            answer=_plain(answer, 50000),
            evidence=uploaded_evidence or [],
        )
    else:
        data = _pptx_bytes(
            title=clean_title,
            question=_plain(question, 12000),
            answer=_plain(answer, 50000),
            evidence=uploaded_evidence or [],
        )
    return data, content_type, _filename(clean_title, extension)


def create_workspace_artifact(
    db: Session,
    *,
    tenant_id: str,
    workspace_id: str | None,
    format_name: str,
    title: str,
    question: str,
    answer: str,
    uploaded_evidence: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    data, content_type, filename = build_artifact_bytes(
        format_name=format_name,
        title=title,
        question=question,
        answer=answer,
        uploaded_evidence=uploaded_evidence or [],
        tenant_id=tenant_id,
    )
    digest = hashlib.sha256(data).hexdigest()
    storage_path: str | None = None
    body_text: str | None = None
    storage_backend = "database_fallback"
    store = None

    if object_storage_configured():
        store = get_object_store()
        with tempfile.NamedTemporaryFile(suffix=Path(filename).suffix, delete=False) as handle:
            handle.write(data)
            temp_path = handle.name
        try:
            stored = store.put_path(
                temp_path,
                tenant_id=tenant_id,
                connection_id=_ARTIFACT_SCOPE,
                filename=filename,
                content_type=content_type,
                expected_sha256=digest,
                expected_size=len(data),
                pending_registration=True,
            )
            storage_path = stored.uri
            storage_backend = "object_storage"
        finally:
            try:
                os.unlink(temp_path)
            except OSError:
                pass
    else:
        body_text = base64.b64encode(data).decode("ascii")

    row = GeneratedArtifact(
        tenant_id=tenant_id,
        workspace_id=workspace_id,
        artifact_type=f"agentic_{str(format_name).lower()}",
        title=_plain(title, 220) or "AGRO-AI Workspace Artifact",
        filename=filename,
        content_type=content_type,
        storage_path=storage_path,
        body_text=body_text,
        metadata_json={
            "source": "ask_agro_ai_agent",
            "sha256": digest,
            "size_bytes": len(data),
            "storage_backend": storage_backend,
            "storage_scope": _ARTIFACT_SCOPE,
            "encoding": "base64" if body_text else None,
            "format": str(format_name).lower(),
        },
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    if store is not None and storage_path:
        store.promote(storage_path, tenant_id=tenant_id, connection_id=_ARTIFACT_SCOPE)

    return {
        "id": row.id,
        "title": row.title,
        "filename": row.filename,
        "content_type": row.content_type,
        "artifact_type": row.artifact_type,
        "workspace_id": row.workspace_id,
        "download_url": f"/v1/agents/artifacts/{row.id}/download",
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


def read_workspace_artifact_bytes(row: GeneratedArtifact) -> bytes:
    metadata = row.metadata_json or {}
    if row.storage_path:
        store = get_object_store()
        return store.read_bytes(
            row.storage_path,
            max_bytes=max(1, int(metadata.get("size_bytes") or 25 * 1024 * 1024)),
            tenant_id=row.tenant_id,
            connection_id=str(metadata.get("storage_scope") or _ARTIFACT_SCOPE),
        )
    if row.body_text and metadata.get("encoding") == "base64":
        return base64.b64decode(row.body_text.encode("ascii"), validate=True)
    if row.body_text:
        return row.body_text.encode("utf-8")
    raise RuntimeError("Artifact content is unavailable")
